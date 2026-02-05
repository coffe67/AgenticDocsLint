from __future__ import annotations

import os
from typing import Any, Dict, List, Optional, Tuple

from slideforge.models import Deck, Slide


def summarize_jira(jira: Dict[str, Any]) -> Dict[str, Any]:
    issues = jira.get("issues", []) or []
    status_counts: Dict[str, int] = {}
    sp_total = 0.0
    sp_done = 0.0
    for it in issues:
        st = str(it.get("status", "Unknown"))
        status_counts[st] = status_counts.get(st, 0) + 1
        sp = it.get("story_points") or 0
        try:
            sp = float(sp)
        except Exception:
            sp = 0.0
        sp_total += sp
        if st.lower() in ("done", "closed", "resolved"):
            sp_done += sp
    return {
        "total_issues": len(issues),
        "status_counts": status_counts,
        "story_points_total": sp_total,
        "story_points_done": sp_done,
    }


def build_deck_from_jira(jira: Dict[str, Any], theme: Optional[Dict[str, Any]] = None) -> Deck:
    theme = theme or {}
    sprint = jira.get("sprint", {}) or {}
    name = sprint.get("name") or "Sprint"
    start = sprint.get("start_date") or ""
    end = sprint.get("end_date") or ""
    summary = summarize_jira(jira)

    slides: List[Slide] = []

    # Title slide
    title_line = f"{name}"
    subtitle = f"{start} – {end}".strip(" – ")
    bullets = [subtitle] if subtitle else []
    slides.append(Slide(index=0, title=title_line, bullets=bullets, notes=None, section="Overview", section_confidence=0.9))

    # Summary slide
    bullets = [
        f"Issues: {summary['total_issues']}",
        f"Story Points (done/total): {summary['story_points_done']:.0f}/{summary['story_points_total']:.0f}",
    ]
    for st, cnt in sorted(summary["status_counts"].items(), key=lambda x: x[0].lower()):
        bullets.append(f"{st}: {cnt}")
    slides.append(Slide(index=1, title="Summary", bullets=bullets, section="Summary", section_confidence=0.9))

    # Top issues slide(s)
    issues = jira.get("issues", []) or []
    top_issues = _select_top_issues(issues, limit=10)
    bullets = [f"{it.get('key')}: {it.get('summary')} [{it.get('status','')}] ({it.get('story_points','?')} SP)" for it in top_issues]
    slides.append(Slide(index=2, title="Top Issues", bullets=bullets, section="Details", section_confidence=0.8))

    deck_meta = {"source": "jira", "sprint": name, "start": start, "end": end, "theme": theme}
    return Deck(meta=deck_meta, slides=slides)


def _select_top_issues(issues: List[Dict[str, Any]], limit: int = 10) -> List[Dict[str, Any]]:
    # Sort by story_points desc, then status
    def keyfn(it: Dict[str, Any]):
        sp = it.get("story_points") or 0
        try:
            sp = float(sp)
        except Exception:
            sp = 0.0
        status = str(it.get("status", "")).lower()
        done_rank = 0 if status in ("done", "closed", "resolved") else 1
        return (-sp, done_rank)

    return sorted(issues, key=keyfn)[:limit]


def export_pptx(deck: Deck, out_path: str, theme: Optional[Dict[str, Any]] = None) -> None:
    try:
        from pptx import Presentation  # type: ignore
        from pptx.util import Inches, Pt  # type: ignore
        from pptx.enum.text import PP_ALIGN  # type: ignore
    except Exception as e:
        raise RuntimeError("PPTX export requires 'python-pptx'. Install via make deps-pptx.") from e

    prs = Presentation()
    # Theme (basic)
    title_font = theme.get("title_font", {}) if theme else {}
    body_font = theme.get("body_font", {}) if theme else {}

    for s in deck.slides:
        layout = prs.slide_layouts[1 if s.index > 0 else 0]  # title slide for first, title+content afterwards
        slide = prs.slides.add_slide(layout)
        title_shape = slide.shapes.title
        if title_shape is not None:
            title_shape.text = s.title
            try:
                for p in title_shape.text_frame.paragraphs:
                    for r in p.runs:
                        if "size" in title_font:
                            r.font.size = Pt(title_font["size"])  # type: ignore
            except Exception:
                pass
        # bullets
        if s.bullets:
            body = slide.placeholders[1] if len(slide.placeholders) > 1 else None
            if body is not None:
                tf = body.text_frame
                tf.clear()
                for i, b in enumerate(s.bullets):
                    if i == 0:
                        tf.text = b
                    else:
                        p = tf.add_paragraph()
                        p.text = b
                try:
                    for p in tf.paragraphs:
                        for r in p.runs:
                            if "size" in body_font:
                                r.font.size = Pt(body_font["size"])  # type: ignore
                except Exception:
                    pass

    prs.save(out_path)


def export_png_summary(jira: Dict[str, Any], out_path: str, theme: Optional[Dict[str, Any]] = None) -> None:
    """Export a simple PNG summary card using Pillow (optional)."""
    try:
        from PIL import Image, ImageDraw, ImageFont  # type: ignore
    except Exception as e:
        raise RuntimeError("PNG export requires 'Pillow'. Install via pip install pillow or make an equivalent target.") from e

    theme = theme or {}
    bg = theme.get("bg", (240, 242, 247))
    fg = theme.get("fg", (20, 22, 30))
    accent = theme.get("accent", (42, 143, 221))
    W, H = (1200, 675)
    im = Image.new("RGB", (W, H), bg)
    draw = ImageDraw.Draw(im)

    sprint = jira.get("sprint", {}) or {}
    name = sprint.get("name") or "Sprint"
    start = sprint.get("start_date") or ""
    end = sprint.get("end_date") or ""
    summary = summarize_jira(jira)

    # Simple text layout
    pad = 40
    y = pad
    title = f"{name}"
    subtitle = f"{start} – {end}".strip(" – ")
    draw.text((pad, y), title, fill=fg)
    y += 60
    if subtitle:
        draw.text((pad, y), subtitle, fill=accent)
        y += 40
    draw.text((pad, y), f"Issues: {summary['total_issues']}", fill=fg)
    y += 32
    draw.text((pad, y), f"Story Points (done/total): {summary['story_points_done']:.0f}/{summary['story_points_total']:.0f}", fill=fg)
    y += 32
    for st, cnt in sorted(summary["status_counts"].items(), key=lambda x: x[0].lower()):
        draw.text((pad, y), f"{st}: {cnt}", fill=fg)
        y += 28

    im.save(out_path)

